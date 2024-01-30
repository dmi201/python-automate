from pyautocad import Autocad, APoint
from openpyxl import Workbook, load_workbook


wb_template = load_workbook(r'C:\Users\Mihai\Desktop\MyFolder\Sablon.xlsx')
ws_t_at = wb_template['Tabel_Antene']

acad = Autocad()

table_data = ws_t_at.values

table = acad.add_table(table_data)
